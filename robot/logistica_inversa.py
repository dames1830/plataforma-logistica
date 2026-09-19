# -*- coding: utf-8 -*-
"""
LOGÍSTICA INVERSA: lo que las tiendas devuelven al CD
====================================================

Arma y publica los datos del módulo **Logística inversa** de la web (Daniel, 19-sep-2026:
*"crea un módulo principal llamado Logística inversa y ahí pones todo lo que hemos creado"*).
Es la maqueta del 19-sep pasada a robot, con las mismas reglas:

  li_vuelve       LO QUE VUELVE: cada guía T del ASN con sus líneas; la pestaña arma adentro las
                  tarjetas, los cuadros por mes, tienda y antigüedad, y las guías con discrepancia.
  li_retorno      DESPACHADO Y DEVUELTO: de lo despachado a cada tienda desde mayo, cuánto retornó,
                  cruzado despacho por despacho (el árbol mes → tienda → modelo).
  li_retorno_det  el detalle del modal de un modelo (picking, embalaje, despacho y las guías T).
                  Va aparte porque es lo más pesado y solo se baja al abrir el primer modal.
  li_doble_tramo  DOBLE TRAMO: las guías T que salen en el correo de comercial con prioridad
                  DOBLE TRAMO (la tienda manda al CD para que se reenvíe a otra tienda).
  li_produccion   PRODUCCIÓN L.I: cada guía T verificada, quién y a qué hora.

Todas con fecha MASTER: la web siempre lee la última. Ver `master-le-gana-a-la-fecha`.

CORRE DETRÁS DEL ASN, no con hora propia: `asn_resumen.py` lo llama al final, igual que a
`recibido_sin_picar.py`, porque lee los seis archivos del ASN recién bajados. NO PUEDE TUMBAR
AL ASN: si falla, el ASN ya está publicado.

LAS REGLAS QUE VIENEN DE LA MAQUETA (y por qué):
- El web report del ASN REPITE la línea una vez por SKU del ASN: se descarta la fila entera
  repetida (918.335 filas T = 46.006 líneas reales, medido).
- Las horas del ASN vienen en UTC: la recepción y la creación se pasan a hora de Lima (-5 h).
  La 'Fecha de envío' es un día sin hora y no se toca.
- La fecha de la guía es la 'Fecha de envío' (la que muestra el ASN de la web).
- Lo que retornó es lo RECIBIDO si la guía ya se ingresó y lo ANUNCIADO si sigue en tránsito.
- LA CALIDAD NO SEPARA: el dígito del medio del código (-1- buena, -9- baja) es una etiqueta
  del CD. Se compara y se cruza por modelo + talla (Daniel, 19-sep-2026).
- Los meses de "ingresó en" son los de la fecha de recepción (lo en tránsito no entra).
- Despachado = OBLPN con Estado de LPN "Enviado", por la hora de asignación de carga; una
  caja con pre-etiqueta (PRE + dígito en el LPN) no es una caja.

    python logistica_inversa.py                 arma y publica en producción
    python logistica_inversa.py --beta          publica en la base de pruebas
    python logistica_inversa.py --probar        arma y NO publica
    python logistica_inversa.py --salida DIR    escribe los JSON en DIR y NO publica
    --asn DIR / --oblpn DIR / --corte AAAA-MM-DD   otras carpetas u otro día de corte
"""
import collections
import csv
import datetime
import glob
import hashlib
import io
import json
import os
import re
import shutil
import sys
import time
import urllib.request

from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
_ARGV = list(sys.argv)
sys.argv = ['x', '--probar']
import armar_pendiente as A          # la carpeta del OneDrive, los correos, las rutas y el token
sys.argv = _ARGV
import maestro_web

T0 = time.time()


def log(t, nivel='INFO'):
    print('[%s] [%-5s] %s' % (datetime.datetime.now().strftime('%H:%M:%S'), nivel, t), flush=True)


A.log = log                           # lo que escribe armar_pendiente va a ESTE log, no al suyo


def arg(nombre, defecto=None):
    return sys.argv[sys.argv.index(nombre) + 1] if nombre in sys.argv and sys.argv.index(nombre) + 1 < len(sys.argv) else defecto


CARPETA_ASN = arg('--asn') or os.path.join(A.BASE, 'ASN')
CARPETA_OBLPN = arg('--oblpn') or os.path.join(A.BASE, 'OBLPN Embalaje')
SALIDA = arg('--salida')
CORTE = (datetime.datetime.strptime(arg('--corte'), '%Y-%m-%d').date() if arg('--corte')
         else datetime.date.today())
# DESDE MAYO (Daniel, 19-sep-2026: "empecemos de mayo"). El mes anterior solo sirve para absorber
# lo que retorna de SUS despachos, para no cargárselo a mayo; no se muestra como fila.
MES_INICIO = '2026-05'
UTC_A_LIMA = datetime.timedelta(hours=5)
MESES_ASN = []                                       # los meses de los archivos del ASN, los llena leer_asn()

ES_PRE = re.compile(r'PRE\d', re.I)                 # pre-etiqueta: no es una caja
PP = re.compile(r'^\d{7}-\d-\d{5}$')                # prepack: 15 caracteres, los pares en 11-12
COLUMNAS_ASN = ['Número de ASN', 'Fecha de envío', 'Proveedor', 'Número de LPN', 'Cantidad enviada',
                'Cantidad recibida', 'Fecha de creación', 'Estado', 'Artículo', 'Descripción',
                'Fecha de recepción', 'cust_field_1', 'cust_field_2', 'cust_field_3', 'cust_field_4',
                'cust_field_5', 'verified_user']


def meses_entre(a, b):
    """['2026-05', '2026-06', ...] de a hasta b, inclusive."""
    y, m = int(a[:4]), int(a[5:7])
    out = []
    while '%04d-%02d' % (y, m) <= b:
        out.append('%04d-%02d' % (y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def mes_anterior(mes):
    y, m = int(mes[:4]), int(mes[5:7])
    return '%04d-%02d' % ((y - 1, 12) if m == 1 else (y, m - 1))


# ══ 1. LAS TIENDAS Y EL MAESTRO ══════════════════════════════════════════════
def leer_rutas():
    """{tienda: [nombre, zona]} del maestro de rutas (CDG, TIENDA, ZONA). Se copia antes de abrirlo:
    en OneDrive puede estar solo en la nube y openpyxl lo ve como un zip roto."""
    ruta = next((r for r in A.RUTAS_CANDIDATOS if os.path.isfile(r)), None)
    if not ruta:
        log('No se encontró el maestro de rutas: las tiendas salen sin nombre ni zona.', 'AVISO')
        return {}
    copia = os.path.join(AQUI, 'logs', '_rutas_copia_li.xlsx')
    try:
        os.makedirs(os.path.dirname(copia), exist_ok=True)
        shutil.copyfile(ruta, copia)
    except Exception:
        copia = ruta
    wb = load_workbook(copia, read_only=True, data_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    cab = [str(c).strip().upper() if c is not None else '' for c in next(it)]
    iC, iN, iZ = cab.index('CDG'), cab.index('TIENDA'), cab.index('ZONA')
    rut = {}
    for f in it:
        if iC < len(f) and f[iC] is not None:
            rut[str(f[iC]).strip().split('.')[0]] = [str(f[iN] or '').strip(), str(f[iZ] or '').strip()]
    wb.close()
    log('Maestro de rutas: %d tiendas' % len(rut))
    return rut


def leer_maestro():
    """{modelo de 7 dígitos: {G. Gender, Gender RIMS, Coleccion PO, Marcas}} del Maestro de la web."""
    try:
        tabla = maestro_web.filas()
    except maestro_web.MaestroNoDisponible as e:
        log('%s Calzado y no calzado no se pueden separar: no se publica.' % e, 'ERROR')
        sys.exit(1)
    if maestro_web.aviso():
        log(maestro_web.aviso(), 'AVISO')
    cab = [str(c).strip() if c is not None else '' for c in tabla[0]]
    idx = {c: i for i, c in enumerate(cab)}
    campos = ('G. Gender', 'Gender RIMS', 'Coleccion PO', 'Marcas')
    ma = {}
    for f in tabla[1:]:
        c = A.limpio(f[idx['CodArticulo']]) if idx['CodArticulo'] < len(f) else ''
        if c:
            ma[c[:7]] = {k: (str(f[idx[k]]).strip() if k in idx and idx[k] < len(f) and f[idx[k]] is not None else '')
                         for k in campos}
    log('Maestro: %s modelos (%s)' % (format(len(ma), ',d'), maestro_web.descripcion()))
    return ma


# ══ 2. LAS LÍNEAS T DEL ASN ══════════════════════════════════════════════════
def a_lima(v):
    """Una fecha con hora del ASN (UTC) en hora de Lima. Un día sin hora queda igual."""
    if not hasattr(v, 'hour'):
        return v
    if (v.hour, v.minute, v.second, v.microsecond) == (0, 0, 0, 0):
        return v
    return v - UTC_A_LIMA


def leer_asn():
    """Las líneas de los ASN que empiezan con T, sin repetidas, en el orden de COLUMNAS_ASN.
    [0] ASN [1] envío [3] LPN [4] enviada [5] recibida [6] creación [7] estado [8] artículo
    [9] descripción [10] recepción [11] tipo [12] CAMBIO [16] quién verificó."""
    lineas, vistas, repetidas = [], set(), 0
    global MESES_ASN
    archivos = sorted(f for f in os.listdir(CARPETA_ASN) if f.lower().endswith('.xlsx') and not f.startswith('~$'))
    if not archivos:
        log('No hay archivos del ASN en %s' % CARPETA_ASN, 'ERROR')
        sys.exit(1)
    for nombre in archivos:
        t = time.time()
        wb = load_workbook(os.path.join(CARPETA_ASN, nombre), read_only=True, data_only=True)
        it = wb.worksheets[0].iter_rows(values_only=True)
        idx = None
        for f in it:
            txt = [str(c).strip() if c is not None else '' for c in (f or ())]
            if 'Número de ASN' in txt:
                idx = {n: i for i, n in enumerate(txt) if n}
                break
        if idx is None:
            log('%s sin encabezado; se salta' % nombre, 'AVISO')
            wb.close()
            continue
        pos = [idx.get(c) for c in COLUMNAS_ASN]
        iN, n0 = pos[0], len(lineas)
        for f in it:
            if not f or iN >= len(f) or f[iN] is None or not str(f[iN]).strip().startswith('T'):
                continue
            fila = tuple(f[i] if i is not None and i < len(f) else None for i in pos)
            # EL ARCHIVO REPITE LA FILA ENTERA: se guarda su huella (la misma llave que asn_resumen.py)
            h = hashlib.md5(''.join('' if c is None else str(c) for c in fila).encode('utf-8', 'replace')).digest()
            if h in vistas:
                repetidas += 1
                continue
            vistas.add(h)
            fila = list(fila)
            fila[0] = str(fila[0]).strip()
            fila[6], fila[10] = a_lima(fila[6]), a_lima(fila[10])
            lineas.append(fila)
        wb.close()
        log('%s: %s líneas T en %.0f s' % (nombre, format(len(lineas) - n0, ',d'), time.time() - t))
    log('ASN: %s líneas T (%s filas repetidas descartadas)' % (format(len(lineas), ',d'), format(repetidas, ',d')))
    # LOS MESES DEL ASN son los de los archivos (uno por mes de CREACIÓN en el WMS). La fecha de la guía puede ser
    # anterior -127 guías de marzo creadas en abril-: la pantalla las junta en "antes de" para que el total cuadre.
    MESES_ASN = sorted({m.group(1) for m in (re.search(r'(\d{4}-\d{2})', a) for a in archivos) if m})
    return lineas


# ══ 3. LO QUE VUELVE ═════════════════════════════════════════════════════════
ESTADO = {'Verified': 'R', 'Receiving Complete': 'R', 'Receiving Started': 'I', 'In Transit': 'T', 'Cancelled': 'C'}


def lo_que_vuelve(U, MA, RUT):
    def clase(art):
        if art.startswith('9920751'):
            return 1                                   # caja H30
        g = (MA.get(art[:7]) or {}).get('G. Gender')
        return 0 if g == 'Footwear' else 2 if g == 'Non Commercial' else 3
    arts, idx, G = [], {}, collections.OrderedDict()
    for r in sorted(U, key=lambda r: (r[0], str(r[8]))):
        a, art = r[0], str(r[8]).strip()
        if art not in idx:
            idx[art] = len(arts)
            arts.append([art, str(r[9] or '').strip(), clase(art)])
        fg = r[1] if hasattr(r[1], 'year') else r[6]  # la fecha de la GUÍA; la de creación si falta
        x = G.get(a)
        if x is None:
            x = G[a] = {'crea': fg, 'recep': r[10], 'est': collections.Counter(), 'tipo': str(r[11] or ''),
                        'cambio': 0, 'user': collections.Counter(), 'E': [0, 0, 0, 0], 'R': [0, 0, 0, 0], 'L': []}
        if fg and (x['crea'] is None or fg < x['crea']):
            x['crea'] = fg
        if r[10] and (x['recep'] is None or r[10] > x['recep']):
            x['recep'] = r[10]
        x['est'][r[7]] += 1
        if r[12] == 'CAMBIO':
            x['cambio'] = 1
        if r[16]:
            x['user'][str(r[16])] += 1
        e, q = int(r[4] or 0), int(r[5] or 0)
        k = arts[idx[art]][2]
        x['E'][k] += e
        x['R'][k] += q
        x['L'].append([idx[art], e, q])
    filas, tiendas = [], {}
    for a, x in G.items():
        t = '50' + a[1:4]
        tiendas[t] = RUT.get(t) or ['(no está en rutas)', '']
        est = ESTADO.get(max(x['est'], key=x['est'].get), '?')
        filas.append([a, t, x['crea'].strftime('%Y-%m-%d %H:%M') if x['crea'] else '',
                      x['recep'].strftime('%Y-%m-%d %H:%M') if x['recep'] and est in 'RI' else '',
                      est, x['tipo'], x['cambio'], (x['user'].most_common(1) or [('', 0)])[0][0],
                      x['E'], x['R'], x['L']])
    meses = MESES_ASN or sorted({f[2][:7] for f in filas if f[2]})
    log('Lo que vuelve: %s guías, %s artículos, %d tiendas' % (format(len(filas), ',d'), format(len(arts), ',d'), len(tiendas)))
    return {'corte': CORTE.isoformat(), 'meses': meses, 'asn': filas, 'tiendas': tiendas, 'arts': arts}


# ══ 4. DESPACHADO Y DEVUELTO ═════════════════════════════════════════════════
def _v(x):
    x = x.strip()
    return x[2:-1] if x.startswith('="') and x.endswith('"') else x


def _num(x):
    try:
        return float(_v(x).replace(',', ''))
    except ValueError:
        return 0.0


def _fecha(x):
    m = re.match(r'(\d{1,2})/(\d{1,2})/(\d{4}) (\d{1,2}):(\d{2})', _v(x))
    return datetime.datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)), int(m.group(4)), int(m.group(5))) if m else None


def _leer_csv(p):
    b = open(p, 'rb').read()
    try:
        t = b.decode('utf-8-sig')
    except UnicodeDecodeError:
        t = b.decode('cp1252')
    cabeza = t[:2000].split('\n', 1)[0]
    r = csv.reader(io.StringIO(t), delimiter=';' if cabeza.count(';') >= cabeza.count(',') else ',')
    next(r, None)
    return r


def _dia_de(nombre):
    m = re.search(r'(\d{1,2})-(\d{1,2})\.csv$', nombre)
    return (int(m.group(2)), int(m.group(1))) if m else (0, 0)


def despachado_y_devuelto(U, MA, RUT):
    MESES = meses_entre(MES_INICIO, CORTE.strftime('%Y-%m'))
    PREVIO = mes_anterior(MES_INICIO)
    desde = datetime.datetime(int(PREVIO[:4]), int(PREVIO[5:7]), 1)
    calzado = lambda s: (MA.get(s[:7]) or {}).get('G. Gender') == 'Footwear'
    pares = lambda s, q: q * int(s[10:12]) if PP.match(s) else q
    DESC = {}
    # ── embalado y despachado: el OBLPN repite la línea en cada archivo del día en que se tocó; se queda
    #    UNA por LPN + artículo + hora de empaquetado, la del archivo más nuevo (trae el último estado) ──
    L = {}
    archivos = sorted(glob.glob(os.path.join(CARPETA_OBLPN, 'OBLPN *.csv')), key=lambda p: (_dia_de(p), os.path.getmtime(p)))
    for p in archivos:
        for r in _leer_csv(p):
            if len(r) < 44:
                continue
            lpn, t, sku = _v(r[3]), _v(r[1]), _v(r[6])
            if not lpn or ES_PRE.search(lpn) or t not in RUT or not calzado(sku):
                continue
            he = _v(r[17])
            if not he:
                continue
            L[(lpn, sku, he)] = (t, pares(sku, _num(r[8])), he, _v(r[5]), _v(r[43]) or _v(r[15]),
                                 _v(r[2]), _v(r[11]), _v(r[12]), _v(r[18]))
            if sku not in DESC:
                DESC[sku] = _v(r[7])
    log('OBLPN: %d archivos, %s líneas únicas de calzado a tienda' % (len(archivos), format(len(L), ',d')))
    des = collections.Counter()
    ENV = collections.defaultdict(list)
    for (lpn, sku, _), (t, q, he, est, hc, orden, emb_u, pick_u, pick_h) in L.items():
        f = _fecha(he)
        if est == 'Enviado':
            fd = _fecha(hc) or f
            if fd and fd >= desde:
                des[(fd.strftime('%Y-%m'), t, sku)] += q
                ENV[(t, sku)].append([fd, q, orden, lpn, he, emb_u, pick_u, pick_h])
    del L
    # ── lo que retornó, casado con el despacho: lo más viejo primero ──
    DEV = collections.defaultdict(list)
    for r in U:
        t, sku = '50' + r[0][1:4], str(r[8]).strip()
        if t not in RUT or not calzado(sku) or r[7] == 'Cancelled':
            continue
        en_transito = r[7] == 'In Transit'
        q_real = float(r[4] or 0) if en_transito else float(r[5] or 0)
        f_guia = r[1] if hasattr(r[1], 'year') else r[6]
        if not q_real or not f_guia:
            continue
        DEV[(t, sku)].append([f_guia, q_real, float(r[5] or 0), r[0], en_transito, r[10]])
        if sku not in DESC:
            DESC[sku] = str(r[9] or '')
    for k in ENV:
        ENV[k].sort()

    def sin_calidad(s):
        return s[:7] + s[9:] if len(s) == 12 and s[7] == '-' and s[9] == '-' else s
    SUEL = collections.defaultdict(list)                   # (tienda, modelo+talla) -> despachos de cualquier calidad
    POOL = collections.defaultdict(list)                   # prepack por (tienda, modelo)
    for (t, sku), l in ENV.items():
        SUEL[(t, sin_calidad(sku))].extend([[x[0], x[1], sku, i] for i, x in enumerate(l)])
        if PP.match(sku):
            POOL[(t, sku[:7])].extend([[x[0], x[1], sku, i] for i, x in enumerate(l)])
    for d in (SUEL, POOL):
        for k in d:
            d[k].sort(key=lambda x: x[0])
    quedanS = {k: [x[1] for x in l] for k, l in SUEL.items()}
    quedanP = {k: [x[1] for x in l] for k, l in POOL.items()}
    casado = collections.defaultdict(lambda: [0.0, collections.Counter(), collections.Counter(), 0.0, collections.Counter()])
    sin = collections.defaultdict(float)
    CAS = collections.defaultdict(list)
    for (t, sku), l in DEV.items():
        for fdev, q, qrec, asn, transito, fing in sorted(l):
            falta = q
            for lista, resto in ((SUEL.get((t, sin_calidad(sku)), []), quedanS.get((t, sin_calidad(sku)))),
                                 (POOL.get((t, sku[:7]), []), quedanP.get((t, sku[:7])))):
                for i, x in enumerate(lista):
                    if falta <= 0:
                        break
                    if x[0] > fdev or resto[i] <= 0:
                        continue
                    usa = min(resto[i], falta)
                    resto[i] -= usa
                    falta -= usa
                    skd, j = x[2], x[3]
                    CAS[(t, skd, j)].append([asn, usa, fdev, transito, sku, fing])
                    k = (x[0].strftime('%Y-%m'), t, skd)
                    casado[k][0] += usa
                    casado[k][1][max(0, (fdev - x[0]).days)] += usa
                    casado[k][2][fdev.strftime('%Y-%m')] += usa
                    if transito:
                        casado[k][3] += usa
                    elif fing:
                        casado[k][4][max(0, (fing - fdev).days)] += usa
            if falta > 0:
                sin[(fdev.strftime('%Y-%m'), t, sku)] += falta
    log('Cruce: %s pares casados con su despacho, %s sin despacho previo'
        % (format(round(sum(x[0] for x in casado.values())), ',d'), format(round(sum(sin.values())), ',d')))

    # ── el árbol: mes de despacho -> tienda -> modelo -> SKU ──
    limpio = lambda d: re.sub(r'(-\d+)+(\.\d)?\s*$', '', d or '').strip(' -')
    col = lambda mod: (MA.get(mod) or {}).get('Coleccion PO') or ''

    def talla(s):
        if PP.match(s):
            return 'prepack %s pares' % int(s[10:12])
        m = re.search(r'-(\d+(?:\.\d)?)\s*$', DESC.get(s, ''))
        return m.group(1) if m else s[-2:]
    # LOS MESES DE "INGRESÓ EN" SON LOS DEL INGRESO: solo lo ingresado, en el mes de su recepción. Si el WMS no
    # trae la fecha y la guía es del último mes, se ingresó en ese mes: no hay otro posible antes del corte.
    ING = collections.defaultdict(collections.Counter)
    for (t, skd, j), lista in CAS.items():
        md = ENV[(t, skd)][j][0].strftime('%Y-%m')
        for a_, usa, fdev, tr_, sr, fi in lista:
            if tr_:
                continue
            mi = fi.strftime('%Y-%m') if fi else (fdev.strftime('%Y-%m') if fdev.strftime('%Y-%m') == MESES[-1] else None)
            if mi in MESES:
                ING[(md, t, skd)][mi] += usa
    arbol = collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(dict)))
    for (m, t, s), q in des.items():
        if m in MESES:
            arbol[m][t][s[:7]][s] = [q, 0.0, {}, [0.0] * len(MESES), 0.0, {}]
    previo = collections.Counter()
    for (m, t, s), (q, dias, pm, tr, ding) in casado.items():
        if m not in MESES:
            for md, p in pm.items():
                if md in MESES:
                    previo[md] += p
            continue
        x = arbol[m][t][s[:7]].setdefault(s, [0.0, 0.0, {}, [0.0] * len(MESES), 0.0, {}])
        x[1] += q
        x[4] += tr
        for d, p in ding.items():
            x[5][str(d)] = x[5].get(str(d), 0) + p
        for d, p in dias.items():
            x[2][str(d)] = x[2].get(str(d), 0) + p
        for mi, p in ING[(m, t, s)].items():
            x[3][MESES.index(mi)] += p
    meses = []
    for m in MESES:
        tiendas = []
        for t, mods in arbol[m].items():
            conDev, resto = [], 0.0
            for mod, skus in mods.items():
                if sum(v[1] for v in skus.values()) > 0:
                    d0 = next((DESC.get(s) for s in skus if DESC.get(s)), '')
                    conDev.append([mod, limpio(d0), col(mod),
                                   sorted([[s, talla(s), round(v[0]), round(v[1]), v[2], [round(p) for p in v[3]], round(v[4]), v[5]]
                                           for s, v in skus.items()])])
                else:
                    resto += sum(v[0] for v in skus.values())
            tiendas.append([t, round(resto), conDev])
        meses.append([m, tiendas])
    sinA = collections.defaultdict(lambda: collections.defaultdict(collections.Counter))
    sinD = {}
    for (m, t, s), q in sin.items():
        if m in MESES:
            sinA[m][t][s[:7]] += q
            if s[:7] not in sinD:
                sinD[s[:7]] = [limpio(DESC.get(s, '')), col(s[:7])]
    sinL = [[m, [[t, [[mod, round(q)] for mod, q in c.most_common()]] for t, c in sinA[m].items()]] for m in MESES]
    # ── el detalle del modal: cada despacho del modelo a la tienda en el mes, con las guías T casadas ──
    corta = lambda x: (x[:5] + ' ' + x[11:16]) if isinstance(x, str) and len(x) >= 16 else (x.strftime('%d/%m %H:%M') if hasattr(x, 'strftime') else '')
    porTM = collections.defaultdict(list)
    for (t, s) in ENV:
        porTM[(t, s[:7])].append(s)
    DET = {}
    for m, tiendas in meses:
        for t, resto, mods in tiendas:
            for mod, desc, colec, skus in mods:
                filas = []
                for s in porTM.get((t, mod), []):
                    for i, x in enumerate(ENV[(t, s)]):
                        if x[0].strftime('%Y-%m') != m:
                            continue
                        g = [[a, round(q, 2), f.strftime('%d/%m/%Y'), 1 if tr else 0, talla(sr), sr, fi.strftime('%d/%m') if fi else '']
                             for a, q, f, tr, sr, fi in CAS.get((t, s, i), [])]
                        filas.append([s, talla(s), corta(x[0]), x[2], x[3], corta(x[7]), x[6], corta(x[4]), x[5], round(x[1]),
                                      sum(y[1] for y in g), g])
                filas.sort(key=lambda r: (r[2][3:5] + r[2][:2] + r[2][6:], r[0]))
                DET[m + '|' + t + '|' + mod] = filas
    tot = lambda i: sum(v[i] for mm in arbol.values() for mods in mm.values() for skus in mods.values() for v in skus.values())
    log('Despachado y devuelto: despachado %s, retornó de eso %s, del mes previo %s, sin despacho %s'
        % tuple(format(round(x), ',d') for x in (tot(0), tot(1), sum(previo.values()),
                                                    sum(q for (m, t, s), q in sin.items() if m in MESES))))
    tiendas = RUT                                          # solo cruzan tiendas del maestro de rutas
    arbolDatos = {'corte': CORTE.isoformat(), 'meses': MESES, 'previo': PREVIO, 'arbol': meses, 'sin': sinL, 'sinDesc': sinD,
                  'abril': [round(previo[m]) for m in MESES], 'tiendas': tiendas}
    return arbolDatos, {'corte': CORTE.isoformat(), 'det': DET}


# ══ 5. DOBLE TRAMO ═══════════════════════════════════════════════════════════
def doble_tramo(RUT):
    """Las guías T del correo de comercial con PRIORIDAD = DOBLE TRAMO. GUIA es la guía T de la tienda que
    manda (T + 3 dígitos de la tienda), TIEND la que recibe y 'Suma de CANTI' las unidades. El correo no trae
    el SKU. Una guía que se repite cuenta una vez, con el PRIMER correo en que aparece.

    Se leen aparte: `armar_pendiente.leer_correos()` deja FUERA justo estas filas (el doble tramo no es
    despacho del CD). De ahí se usa lo mismo que ya resolvió: la hoja que trae la columna GUIA y las
    columnas por nombre, porque la cabecera del correo cambió seis veces."""
    if not os.path.isdir(A.CORREOS):
        raise RuntimeError('no existe la carpeta de correos %s' % A.CORREOS)
    filas, vistas = [], set()
    for (mes, dia), nombre in A.archivos_de_correo():
        try:
            wb = load_workbook(os.path.join(A.CORREOS, nombre), read_only=True, data_only=True)
        except Exception as e:
            log('No se pudo abrir %s (%s)' % (nombre, type(e).__name__), 'AVISO')
            continue
        anio = CORTE.year if mes <= CORTE.month else CORTE.year - 1     # el nombre del correo no trae el año
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            try:
                cols = A.columnas_del_correo([str(c).strip() if c is not None else '' for c in next(it)])
            except StopIteration:
                continue
            if A.C_GUIA not in cols or A.C_CANT not in cols:
                continue
            for r in it:
                f = A.fila_del_correo(r, cols)
                if 'DOBLE' not in str(f[A.C_PRIOR] or '').upper():
                    continue
                g = A.limpio(f[A.C_GUIA])
                if not re.match(r'^T\d{11}$', g) or g in vistas:
                    continue
                vistas.add(g)
                try:
                    q = float(f[A.C_CANT] or 0)
                except (TypeError, ValueError):
                    q = 0
                d = str(f[A.C_TIEND] or '').strip().split('.')[0]      # el correo dice 238; la tienda es 50238
                filas.append([g, '50' + g[1:4], d if len(d) == 5 and d.startswith('50') else '50' + d.zfill(3),
                              '%04d-%02d' % (anio, mes), round(q), str(f[A.C_ETIQ] or '').strip()])
            break
        wb.close()
    tiendas = {t: RUT.get(t, ['(no está en rutas)', '']) for x in filas for t in (x[1], x[2])}
    log('Doble tramo: %s guías, %s unidades' % (format(len(filas), ',d'), format(sum(x[4] for x in filas), ',d')))
    return {'corte': CORTE.isoformat(), 'meses': sorted({x[3] for x in filas}), 'guias': filas, 'tiendas': tiendas}


# ══ 6. PRODUCCIÓN L.I ════════════════════════════════════════════════════════
def produccion(U, MA):
    """Cada guía T verificada es un evento: quién la verificó, a qué hora (el WMS guarda UNA hora por guía)
    y cuántas unidades se recibieron. Calzado = G. Gender Footwear del Maestro."""
    ev = collections.OrderedDict()
    res = collections.defaultdict(lambda: {'col': collections.Counter(), 'mar': collections.Counter(), 'gen': collections.Counter(),
                                           'colc': collections.Counter(), 'marc': collections.Counter(), 'genc': collections.Counter()})
    for r in sorted((r for r in U if r[10]), key=lambda r: (r[10], r[0])):
        if not r[10] or not r[16] or not r[5]:
            continue
        q = float(r[5])
        m = MA.get(str(r[8])[:7]) or {}
        cal = m.get('G. Gender') == 'Footwear'
        e = ev.setdefault(r[0], [r[0], str(r[16]), r[10].strftime('%Y-%m-%d %H:%M:%S'), 0, 0, 0])
        e[3 if cal else 4] += q
        e[5] += 1
        for mes in ('Todo', r[10].strftime('%Y-%m')):
            x = res[mes]
            for k, v in (('col', m.get('Coleccion PO') or 'Sin dato'), ('mar', m.get('Marcas') or 'Sin dato'),
                         ('gen', m.get('Gender RIMS') or 'Sin dato')):
                x[k][v] += q
                if cal:
                    x[k + 'c'][v] += q
    out = {'corte': CORTE.isoformat(), 'meses': sorted(m for m in res if m != 'Todo'),
           'ev': [[e[0], e[1], e[2], round(e[3]), round(e[4]), e[5]] for e in ev.values()],
           'res': {mes: {k: [[n, round(q), round(x[k + 'c'][n])] for n, q in x[k].most_common()] for k in ('col', 'mar', 'gen')}
                   for mes, x in res.items()}}
    log('Producción L.I: %s guías verificadas por %d personas' % (format(len(out['ev']), ',d'), len({e[1] for e in out['ev']})))
    return out


# ══ 7. PUBLICAR ══════════════════════════════════════════════════════════════
def publicar(area, datos):
    cuerpo = json.dumps(datos, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
    if SALIDA:
        os.makedirs(SALIDA, exist_ok=True)
        with open(os.path.join(SALIDA, area + '.json'), 'wb') as fh:
            fh.write(cuerpo)
        log('%s: %.0f KB escritos en %s (no se publica)' % (area, len(cuerpo) / 1024.0, SALIDA))
        return True
    if '--probar' in sys.argv:
        log('MODO PROBAR: %s no se publica (serían %.0f KB)' % (area, len(cuerpo) / 1024.0))
        return True
    url = '%s/%s?date=MASTER' % (A.WEB_DATOS_API, area)
    for intento in range(1, 4):
        try:
            p = urllib.request.Request(url, data=cuerpo, method='POST')
            p.add_header('Content-Type', 'application/json')
            p.add_header('X-Robot-Token', A.ROBOT_TOKEN)
            if '--beta' in sys.argv:
                p.add_header('X-Environment', 'beta')
            with urllib.request.urlopen(p, timeout=300) as resp:
                json.loads(resp.read().decode('utf-8'))
            log('Publicado %s: %.0f KB' % (area, len(cuerpo) / 1024.0))
            return True
        except Exception as e:
            if intento < 3:
                log('Intento %d: no se pudo publicar %s (%s), se reintenta' % (intento, area, type(e).__name__), 'AVISO')
                time.sleep(20)
            else:
                log('No se pudo publicar %s: %s: %s' % (area, type(e).__name__, str(e)[:160]), 'ERROR')
    return False


def main():
    log('=' * 58)
    log('LOGÍSTICA INVERSA PARA LA WEB (corte %s)' % CORTE.isoformat())
    log('=' * 58)
    RUT = leer_rutas()
    MA = leer_maestro()
    U = leer_asn()
    ok = True
    ok &= publicar('li_vuelve', lo_que_vuelve(U, MA, RUT))
    ok &= publicar('li_produccion', produccion(U, MA))
    try:
        ok &= publicar('li_doble_tramo', doble_tramo(RUT))
    except Exception as e:                   # sin correos el resto se publica igual
        log('Doble tramo: no se pudieron leer los correos (%s: %s)' % (type(e).__name__, str(e)[:160]), 'ERROR')
        ok = False
    arbol, det = despachado_y_devuelto(U, MA, RUT)
    # el detalle PRIMERO: la web abre el árbol y pide el detalle recién al hacer clic en un modelo
    ok &= publicar('li_retorno_det', det)
    ok &= publicar('li_retorno', arbol)
    log('Listo en %.0f s%s' % (time.time() - T0, '' if ok else ' CON ERRORES'))
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
