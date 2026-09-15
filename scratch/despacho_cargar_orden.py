# -*- coding: utf-8 -*-
"""
DEJAR BETA CON UNA SOLA ORDEN DE DESPACHO

Daniel, 15-sep-2026: *"borres todo y solamente te quedes con ese archivo del 15 del 9,
el que te acabo de pasar; con ese hay que hacer pruebas"*.

Borra las areas del despacho en BETA y carga unicamente las guias del archivo que manda
comercial por WhatsApp. Produccion NO SE TOCA: este script no acepta --produccion.

Antes de borrar se guarda una copia local de todo lo que habia. La historia igual no se
pierde -las 3.129 guias siguen en produccion y en el AppSheet-, pero el respaldo cuesta
un segundo y evita depender de eso.

    python cargar_od.py "C:\\ruta\\Despacho del 15.09.xlsx"            # prueba en seco
    python cargar_od.py "C:\\ruta\\Despacho del 15.09.xlsx" --subir
"""
import sys
import io
import os
import json
import re
import datetime
import urllib.request
import openpyxl

API = 'https://logistics-backend-wv0x.onrender.com/api/logistics'
BETA = {'X-Environment': 'beta'}

AREA_BASE = 'despacho_catalogo'
AREA_INDICE = 'despacho_catalogo_indice'
AREA_CAMBIOS = 'despacho_catalogo_cambios'
AREA_SEMANA = 'despacho_catalogo_sem_'

# Los 14 titulos que manda comercial, en el orden en que vienen. Son los mismos que las
# columnas E a R de la BBDD, con los mismos nombres: por eso se reconocen por titulo y
# no por posicion, que es lo unico que sobrevive a que alguien mueva una columna.
TITULOS = {
    'asesor': 'ase', 'NombreLider': 'lider', 'Promotor': 'prom', 'Rotulo': 'rot',
    'Agencia': 'age', 'Destino': 'dest', 'Pedido': 'ped', 'Total_Cantidad': 'cant',
    'PedidoBolsas': 'pedBol', 'TotalBolsas': 'bolsas', 'TotalVenta': 'venta',
    'CobroFlete': 'flete', 'Observacion': 'obs', 'Detalle': 'det'
}
NUMEROS = ('cant', 'bolsas', 'venta', 'pedBol')
CAMPOS_CATALOGO = ('ase', 'lider', 'age', 'dest', 'est', 'flete', 'obs')


def limpio(v):
    """Los titulos vienen con un espacio duro pegado a los lados. Sin sacarlo no
    coincide NI UNO de los catorce."""
    return re.sub(r'\s+', ' ', str(v if v is not None else '').replace(u'\xa0', ' ')).strip()


def pedir(area):
    pet = urllib.request.Request('%s/%s?date=MASTER' % (API, area), headers=BETA)
    with urllib.request.urlopen(pet, timeout=180) as r:
        c = json.loads(r.read().decode('utf-8'))
    return c.get('data', c) if isinstance(c, dict) else c


def guardar(area, datos):
    cuerpo = json.dumps(datos, ensure_ascii=False).encode('utf-8')
    cab = {'Content-Type': 'application/json'}
    cab.update(BETA)
    pet = urllib.request.Request('%s/%s?date=MASTER' % (API, area), data=cuerpo,
                                 headers=cab, method='POST')
    with urllib.request.urlopen(pet, timeout=180) as r:
        r.read()
    return len(cuerpo)


def leer_orden(ruta):
    """Devuelve (cabecera, filas). Busca la hoja y la fila de titulos: ninguna de las dos
    esta en un sitio fijo."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    for ws in wb.worksheets:
        # LA FILA DE TITULOS SE BUSCA, no se fija en la 9: arriba hay un resumen que hoy
        # ocupa 8 filas y manana puede ocupar 7 o 10.
        for r in range(1, min(ws.max_row, 40) + 1):
            fila = [limpio(ws.cell(row=r, column=c).value)
                    for c in range(1, ws.max_column + 1)]
            if sum(1 for t in TITULOS if t in fila) >= 10:
                cab = {}
                for c in range(1, ws.max_column + 1):
                    t = limpio(ws.cell(row=r, column=c).value)
                    if t in TITULOS:
                        cab[TITULOS[t]] = c
                # La cabecera de arriba: numero de orden y fecha, buscados por su rotulo.
                meta = {}
                for rr in range(1, r):
                    for cc in range(1, ws.max_column):
                        et = limpio(ws.cell(row=rr, column=cc).value).lower().rstrip('.')
                        if et in ('nro documento', 'nro. documento', 'n documento'):
                            meta['od'] = limpio(ws.cell(row=rr, column=cc + 1).value)
                        elif et.startswith('fec. creaci') or et.startswith('fec creaci'):
                            meta['creada'] = str(ws.cell(row=rr, column=cc + 1).value or '')[:10]
                        elif et == 'pares enviado':
                            meta['paresEnviado'] = ws.cell(row=rr, column=cc + 1).value
                        elif et == 'pares pedido':
                            meta['paresPedido'] = ws.cell(row=rr, column=cc + 1).value
                        elif et == 'total monto':
                            meta['monto'] = ws.cell(row=rr, column=cc + 1).value
                filas = []
                for rr in range(r + 1, ws.max_row + 1):
                    g = {}
                    for campo, c in cab.items():
                        v = ws.cell(row=rr, column=c).value
                        if v is None or limpio(v) == '':
                            continue
                        g[campo] = float(v) if (campo in NUMEROS and isinstance(v, (int, float))) else limpio(v)
                    # Una fila sin rotulo y sin pedido esta vacia: es el relleno del final.
                    if g.get('rot') or g.get('ped'):
                        filas.append(g)
                meta['hoja'] = ws.title
                return meta, filas
    raise SystemExit('No encontre la fila de titulos en ninguna hoja de %s' % ruta)


def compactar(filas):
    cat = {}
    for campo in CAMPOS_CATALOGO:
        vistos, pos = [], {}
        for f in filas:
            v = f.get(campo)
            if isinstance(v, str) and v and v not in pos:
                pos[v] = len(vistos)
                vistos.append(v)
        if len(vistos) > 1:
            cat[campo] = vistos
    apretadas = []
    for f in filas:
        g = {}
        for k, v in f.items():
            g[k] = cat[k].index(v) if (k in cat and isinstance(v, str)) else v
        apretadas.append(g)
    return cat, apretadas


def lunes_de(t):
    d = datetime.date(int(t[0:4]), int(t[5:7]), int(t[8:10]))
    return (d - datetime.timedelta(days=d.weekday())).isoformat()


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not args:
        print(__doc__)
        return 1
    ruta = args[0]
    subir = '--subir' in sys.argv
    hoy = datetime.date.today().isoformat()
    # LA FECHA DE DESPACHO ES LA DEL DIA QUE SALE, no la de creacion del documento.
    # Daniel: *"ese es el archivo que va a salir hoy dia de despacho 15"*. El archivo se
    # llama "del 15.09" y adentro dice creado el 14: el 14 es cuando comercial armo la
    # orden, el 15 es cuando sale el camion.
    fecha = next((a.split('=')[1] for a in sys.argv if a.startswith('--fecha=')), hoy)

    meta, filas = leer_orden(ruta)
    print('ARCHIVO   %s' % os.path.basename(ruta))
    print('  hoja    %s' % meta.get('hoja'))
    print('  orden   %s   creada el %s' % (meta.get('od', '?'), meta.get('creada', '?')))
    print('  guias   %d' % len(filas))
    print('  FECHA DE DESPACHO QUE SE LES PONE: %s' % fecha)

    pares = sum(float(f.get('cant') or 0) for f in filas)
    venta = sum(float(f.get('venta') or 0) for f in filas)
    print('\nCUADRE CONTRA EL RESUMEN DEL PROPIO ARCHIVO:')
    print('  pares  %8.0f   el archivo dice Pares Pedido %s' % (pares, meta.get('paresPedido')))
    print('  venta  %8.2f   el archivo dice Total Monto  %s' % (venta, meta.get('monto')))
    cuadra = (str(meta.get('paresPedido')) == str(int(pares))
              and abs(float(meta.get('monto') or 0) - venta) < 0.05)
    print('  %s' % ('CUADRA' if cuadra else 'NO CUADRA - revisar antes de cargar'))

    # Numeracion: se sigue el correlativo del AppSheet para no chocar si algun dia se
    # vuelve a importar de alla.
    try:
        base = pedir(AREA_BASE) or {}
        vistos = [int(x['id']) for x in (base.get('filas') or []) if str(x.get('id', '')).isdigit()]
        desde = (max(vistos) + 1) if vistos else 40001
        print('\nEN BETA HABIA %d guias. Se van a BORRAR.' % len(base.get('filas') or []))
    except Exception as e:
        desde = 40001
        print('\n(no se pudo leer lo que habia en beta: %s)' % e)

    for i, f in enumerate(filas):
        f['id'] = str(desde + i)
        f['desp'] = fecha
        f['est'] = 'PENDIENTE'
        f['canal'] = 'catalogo'
        f['od'] = meta.get('od', '')
    print('  numeros: del %s al %s' % (filas[0]['id'], filas[-1]['id']))

    if not subir:
        print('\nPrueba en seco. Para hacerlo de verdad: --subir')
        return 0

    # ── RESPALDO ANTES DE BORRAR ─────────────────────────────────────────────
    resp = {}
    for a in [AREA_BASE, AREA_INDICE, AREA_CAMBIOS]:
        try:
            resp[a] = pedir(a)
        except Exception:
            resp[a] = None
    viejo = (resp.get(AREA_INDICE) or {}).get('semanas') or {}
    for l in viejo:
        try:
            resp[AREA_SEMANA + l] = pedir(AREA_SEMANA + l)
        except Exception:
            pass
    nombre = 'respaldo_beta_despacho_%s.json' % datetime.datetime.now().strftime('%Y%m%d_%H%M')
    io.open(nombre, 'w', encoding='utf-8').write(json.dumps(resp, ensure_ascii=False))
    print('\nRespaldo de lo que habia: %s (%d areas)' % (nombre, len(resp)))

    # ── BORRAR ───────────────────────────────────────────────────────────────
    print('\nBorrando en beta...')
    for a in [AREA_BASE, AREA_CAMBIOS] + [AREA_SEMANA + l for l in viejo]:
        guardar(a, {})
        print('   vacia  %s' % a)

    # ── CARGAR ───────────────────────────────────────────────────────────────
    paquete = {'version': 2, 'origen': os.path.basename(ruta), 'od': meta.get('od', ''),
               'importado': hoy, 'filas': filas}
    guardar(AREA_BASE, paquete)
    print('\n   ok  %s con %d guias' % (AREA_BASE, len(filas)))

    lunes = lunes_de(fecha)
    cat, apretadas = compactar(filas)
    guardar(AREA_SEMANA + lunes, {'version': 2, 'lunes': lunes, 'cat': cat, 'filas': apretadas})
    print('   ok  %s%s' % (AREA_SEMANA, lunes))

    indice = {'version': 2, 'total': len(filas), 'desde': fecha, 'hasta': fecha,
              'importado': hoy, 'origen': os.path.basename(ruta),
              'partido': datetime.datetime.now().strftime('%Y-%m-%d %H:%M'),
              'semanas': {lunes: {
                  'n': len(filas), 'sin': len(filas),
                  'sinPorCanal': {'catalogo': len(filas), 'retail': 0},
                  'porCanal': {'catalogo': len(filas), 'retail': 0},
                  'd0': fecha, 'd1': fecha,
                  'kb': int(round(len(json.dumps(apretadas, ensure_ascii=False)) / 1024.0))}}}
    guardar(AREA_INDICE, indice)
    print('   ok  %s' % AREA_INDICE)
    print('\nListo. Beta tiene %d guias, todas en PENDIENTE, del %s.' % (len(filas), fecha))
    return 0


if __name__ == '__main__':
    sys.exit(main())
